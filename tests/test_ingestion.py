from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
import pytest

from submission_evaluation.agents import ingestion as ingestion_module
from submission_evaluation.agents.ingestion import ExtractedLocationCandidate, IngestionAgent
from submission_evaluation.config import AppConfig


def _write_xlsx(path: Path, rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Locations"
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def _excel_rows_for_locations(count: int) -> list[list[object]]:
    rows: list[list[object]] = [["Location", "Address", "City", "State", "Zip", "TIV"]]
    for idx in range(1, count + 1):
        rows.append(
            [
                f"Site {idx}",
                f"{idx} Main St",
                "Austin",
                "TX",
                f"78{idx:03d}",
                float(1_000_000 + idx),
            ]
        )
    return rows


def _build_config(input_dir: Path) -> AppConfig:
    return AppConfig(
        input_dir=input_dir,
        processed_dir=input_dir / "processed",
        output_dir=input_dir / "output",
        azure_openai_endpoint="",
        azure_openai_api_key="",
        azure_openai_deployment="",
        azure_openai_api_version="2024-06-01",
        web_search_endpoint="",
        web_search_api_key="",
    )


class _StubPdfExtractor:
    def __init__(self, candidates: list[dict[str, Any]] | None = None) -> None:
        self._candidates = candidates or []

    def extract_location_candidates(self, _text: str):
        return [
            ExtractedLocationCandidate(
                location_name=candidate.get("location_name"),
                address_line_1=candidate.get("address_line_1"),
                city=candidate.get("city"),
                state=candidate.get("state"),
                postal_code=candidate.get("postal_code"),
                country=candidate.get("country"),
                confidence=float(candidate.get("confidence", 0.8)),
                supplementary_text=candidate.get("supplementary_text"),
                extraction_method="stub",
            )
            for candidate in self._candidates
        ]


def test_parse_locations_when_xlsx_contains_headers_aliases_and_empty_rows_should_extract_locations(
    tmp_path: Path,
):
    input_dir = tmp_path / "input"
    input_dir.mkdir()

    _write_xlsx(
        input_dir / "submission.xlsx",
        [
            ["Submission sheet", "ignore me"],
            ["Location Name", "Street Address", "City", "State", "Zip", "TIV"],
            ["HQ", "100 Main St", "Austin", "TX", "78701", "2500000"],
            ["Warehouse", "200 Market St", "Dallas", "TX", "75201", 1400000],
            [None, None, None, None, None, None],
        ],
    )

    agent = IngestionAgent(config=_build_config(input_dir))
    manifest, warnings = agent.parse_locations(submission_id="sub-100")

    assert warnings == []
    assert manifest.count == 2
    assert manifest.locations[0].location_name == "HQ"
    assert manifest.locations[1].location_name == "Warehouse"
    assert manifest.locations[0].total_insured_value == 2500000.0
    assert manifest.locations[0].source_row == 3

    manifest_b, _ = agent.parse_locations(submission_id="sub-100")
    assert manifest.manifest_hash == manifest_b.manifest_hash
    assert [location.id for location in manifest.locations] == [location.id for location in manifest_b.locations]


def test_parse_locations_when_standard_excel_has_10_locations_should_preserve_exact_count(
    tmp_path: Path,
):
    input_dir = tmp_path / "input"
    input_dir.mkdir()

    _write_xlsx(input_dir / "ten-sites.xlsx", _excel_rows_for_locations(10))

    agent = IngestionAgent(config=_build_config(input_dir))
    manifest, warnings = agent.parse_locations(submission_id="sub-10")

    assert warnings == []
    assert manifest.count == 10
    assert len(manifest.locations) == 10


def test_parse_locations_when_excel_has_100_locations_should_preserve_max_supported_count(
    tmp_path: Path,
):
    input_dir = tmp_path / "input"
    input_dir.mkdir()

    _write_xlsx(input_dir / "hundred-sites.xlsx", _excel_rows_for_locations(100))

    agent = IngestionAgent(config=_build_config(input_dir))
    manifest, warnings = agent.parse_locations(submission_id="sub-100-max")

    assert warnings == []
    assert manifest.count == 100
    assert len({location.id for location in manifest.locations}) == 100


def test_parse_locations_when_address_missing_should_include_location_and_flag_missing_critical_field(
    tmp_path: Path,
):
    input_dir = tmp_path / "input"
    input_dir.mkdir()

    _write_xlsx(
        input_dir / "missing-address.xlsx",
        [
            ["Location", "Address", "City", "State"],
            ["Site A", "", "Austin", "TX"],
        ],
    )

    agent = IngestionAgent(config=_build_config(input_dir))
    manifest, warnings = agent.parse_locations(submission_id="sub-101")

    assert warnings == []
    assert manifest.count == 1
    flagged = manifest.locations[0].additional_fields.get("missing_critical_fields")
    assert flagged == ("address_line_1",)


def test_parse_locations_when_multiple_excel_files_should_aggregate_rows(tmp_path: Path):
    input_dir = tmp_path / "input"
    input_dir.mkdir()

    _write_xlsx(
        input_dir / "one.xlsx",
        [
            ["Location", "Address", "City", "State"],
            ["A", "10 First St", "Austin", "TX"],
        ],
    )
    _write_xlsx(
        input_dir / "two.xlsx",
        [
            ["Location", "Address", "City", "State"],
            ["B", "20 Second St", "Dallas", "TX"],
        ],
    )

    agent = IngestionAgent(config=_build_config(input_dir))
    manifest, _ = agent.parse_locations(submission_id="sub-102")

    assert manifest.count == 2
    assert {location.location_name for location in manifest.locations} == {"A", "B"}
    assert len({location.id for location in manifest.locations}) == 2


def test_parse_locations_when_file_is_corrupt_should_log_warning_and_continue(
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
):
    input_dir = tmp_path / "input"
    input_dir.mkdir()

    _write_xlsx(
        input_dir / "good.xlsx",
        [
            ["Location", "Address", "City", "State"],
            ["Valid", "10 First St", "Austin", "TX"],
        ],
    )
    (input_dir / "bad.xlsx").write_text("not an excel file", encoding="utf-8")

    agent = IngestionAgent(config=_build_config(input_dir))
    with caplog.at_level("WARNING"):
        manifest, warnings = agent.parse_locations(submission_id="sub-103")

    assert manifest.count == 1
    assert len(warnings) == 1
    assert "bad.xlsx" in warnings[0]
    assert "Skipped unreadable file" in caplog.text


def test_parse_locations_when_xls_dependency_missing_should_skip_xls_with_warning(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    input_dir = tmp_path / "input"
    input_dir.mkdir()

    _write_xlsx(
        input_dir / "good.xlsx",
        [
            ["Location", "Address", "City", "State"],
            ["Valid", "10 First St", "Austin", "TX"],
        ],
    )
    (input_dir / "legacy.xls").write_bytes(b"placeholder")

    monkeypatch.setattr(ingestion_module, "xlrd", None)
    agent = IngestionAgent(config=_build_config(input_dir))

    manifest, warnings = agent.parse_locations(submission_id="sub-104")

    assert manifest.count == 1
    assert len(warnings) == 1
    assert "legacy.xls" in warnings[0]
    assert "xlrd" in warnings[0]


def test_parse_locations_when_file_temporarily_locked_should_retry_and_succeed(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    file_path = input_dir / "locked.xlsx"
    _write_xlsx(
        file_path,
        [
            ["Location", "Address", "City", "State"],
            ["Retry Site", "10 Retry St", "Austin", "TX"],
        ],
    )

    attempts = {"count": 0}
    original_load_workbook = ingestion_module.load_workbook

    def fake_load_workbook(*_args, **_kwargs):
        attempts["count"] += 1
        if attempts["count"] < 3:
            raise PermissionError("file locked")
        return original_load_workbook(*_args, **_kwargs)

    monkeypatch.setattr(ingestion_module, "load_workbook", fake_load_workbook)
    monkeypatch.setattr(ingestion_module.time, "sleep", lambda _seconds: None)
    agent = IngestionAgent(config=_build_config(input_dir))

    manifest, warnings = agent.parse_locations(submission_id="sub-105")

    assert warnings == []
    assert manifest.count == 1
    assert attempts["count"] == 3


def test_parse_locations_when_pdf_has_no_locations_should_continue_without_error(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    _write_xlsx(
        input_dir / "submission.xlsx",
        [
            ["Location", "Address", "City", "State"],
            ["Excel Site", "100 Main St", "Austin", "TX"],
        ],
    )
    (input_dir / "supplement.pdf").write_bytes(b"%PDF-1.4 placeholder")

    monkeypatch.setattr(ingestion_module, "_extract_pdf_text_from_file", lambda _p: "narrative only")
    agent = IngestionAgent(config=_build_config(input_dir), pdf_location_extractor=_StubPdfExtractor())

    manifest, warnings = agent.parse_locations(submission_id="sub-200")

    assert warnings == []
    assert manifest.count == 1
    assert manifest.locations[0].address_line_1 == "100 Main St"


def test_parse_locations_when_pdf_location_duplicates_excel_should_dedupe_by_address(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    _write_xlsx(
        input_dir / "submission.xlsx",
        [
            ["Location", "Address", "City", "State", "Zip"],
            ["Excel HQ", "100 Main St", "Austin", "TX", "78701"],
        ],
    )
    (input_dir / "submission.pdf").write_bytes(b"%PDF-1.4 placeholder")

    monkeypatch.setattr(ingestion_module, "_extract_pdf_text_from_file", lambda _p: "location text")
    agent = IngestionAgent(
        config=_build_config(input_dir),
        pdf_location_extractor=_StubPdfExtractor(
            [
                {
                    "location_name": "PDF HQ",
                    "address_line_1": "100 Main St",
                    "city": "Austin",
                    "state": "TX",
                    "postal_code": "78701",
                    "confidence": 0.93,
                    "supplementary_text": "Narrative details",
                }
            ]
        ),
    )

    manifest, _ = agent.parse_locations(submission_id="sub-201")

    assert manifest.count == 1
    location = manifest.locations[0]
    assert "submission.pdf" in location.additional_fields["pdf_sources"]
    assert "Narrative details" in location.additional_fields["pdf_supplementary"]


def test_parse_locations_when_pdf_extraction_low_confidence_should_flag_audit_metadata(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    (input_dir / "submission.pdf").write_bytes(b"%PDF-1.4 placeholder")

    monkeypatch.setattr(ingestion_module, "_extract_pdf_text_from_file", lambda _p: "location text")
    agent = IngestionAgent(
        config=_build_config(input_dir),
        pdf_location_extractor=_StubPdfExtractor(
            [
                {
                    "location_name": "Low Confidence Site",
                    "address_line_1": "22 River Rd",
                    "city": "Austin",
                    "state": "TX",
                    "postal_code": "78702",
                    "confidence": 0.4,
                }
            ]
        ),
    )

    manifest, warnings = agent.parse_locations(submission_id="sub-202")

    assert warnings == []
    assert manifest.count == 1
    location = manifest.locations[0]
    assert "low_confidence_pdf_extraction" in location.additional_fields["audit_flags"]
    assert 0.4 in location.additional_fields["pdf_extraction_confidence"]


def test_parse_locations_when_pdf_has_existing_and_new_locations_should_merge_and_rehash(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    _write_xlsx(
        input_dir / "submission.xlsx",
        [
            ["Location", "Address", "City", "State", "Zip"],
            ["Excel HQ", "100 Main St", "Austin", "TX", "78701"],
        ],
    )
    (input_dir / "locations.pdf").write_bytes(b"%PDF-1.4 placeholder")

    monkeypatch.setattr(ingestion_module, "_extract_pdf_text_from_file", lambda _p: "location text")

    base_agent = IngestionAgent(config=_build_config(input_dir), pdf_location_extractor=_StubPdfExtractor())
    base_manifest, _ = base_agent.parse_locations(submission_id="sub-203")

    merged_agent = IngestionAgent(
        config=_build_config(input_dir),
        pdf_location_extractor=_StubPdfExtractor(
            [
                {
                    "location_name": "PDF Duplicate",
                    "address_line_1": "100 Main St",
                    "city": "Austin",
                    "state": "TX",
                    "postal_code": "78701",
                    "confidence": 0.91,
                    "supplementary_text": "Same location narrative",
                },
                {
                    "location_name": "PDF New Site",
                    "address_line_1": "500 Elm St",
                    "city": "Dallas",
                    "state": "TX",
                    "postal_code": "75201",
                    "confidence": 0.84,
                },
            ]
        ),
    )

    merged_manifest, warnings = merged_agent.parse_locations(submission_id="sub-203")

    assert warnings == []
    assert merged_manifest.count == 2
    assert merged_manifest.manifest_hash != base_manifest.manifest_hash
    added = [location for location in merged_manifest.locations if location.address_line_1 == "500 Elm St"]
    assert len(added) == 1
    assert added[0].source_sheet == "pdf"


def test_parse_locations_when_input_files_provided_should_only_process_selected_batch(tmp_path: Path):
    input_dir = tmp_path / "input"
    input_dir.mkdir()

    selected_file = input_dir / "selected.xlsx"
    ignored_file = input_dir / "ignored.xlsx"
    _write_xlsx(
        selected_file,
        [
            ["Location", "Address", "City", "State"],
            ["Selected Site", "10 First St", "Austin", "TX"],
        ],
    )
    _write_xlsx(
        ignored_file,
        [
            ["Location", "Address", "City", "State"],
            ["Ignored Site", "20 Second St", "Dallas", "TX"],
        ],
    )

    agent = IngestionAgent(config=_build_config(input_dir))

    manifest, warnings = agent.parse_locations(
        submission_id="sub-204",
        input_files=[str(selected_file)],
    )

    assert warnings == []
    assert manifest.count == 1
    assert manifest.locations[0].location_name == "Selected Site"
