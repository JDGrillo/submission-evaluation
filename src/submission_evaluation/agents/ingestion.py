from __future__ import annotations

from dataclasses import dataclass, fields, replace
from pathlib import Path
import logging
import re
import time
from typing import Any, Mapping, Protocol, Sequence
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ..clients import AzureOpenAIClient
from ..config import AppConfig
from ..models import Location, LocationManifest
from .base import AgentResult, NoOpAgent

try:
    import xlrd  # type: ignore[import-not-found]
except ModuleNotFoundError:  # pragma: no cover - exercised in tests via behavior
    xlrd = None

try:
    from pypdf import PdfReader  # type: ignore[import-not-found]
except ModuleNotFoundError:  # pragma: no cover - exercised in tests via behavior
    PdfReader = None


LOGGER = logging.getLogger(__name__)

_HEADER_SCAN_LIMIT = 30
_FILE_READ_RETRIES = 3
_RETRY_BACKOFF_SECONDS = 0.1
_CRITICAL_FIELDS = ("address_line_1",)
_LOW_CONFIDENCE_THRESHOLD = 0.7
_INTERNAL_FIELDS = {
    "id",
    "submission_id",
    "source_file",
    "source_sheet",
    "source_row",
    "additional_fields",
}

_FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "location_name": ("location", "location name", "site name", "name"),
    "address_line_1": (
        "address",
        "address 1",
        "address line 1",
        "street",
        "street address",
        "location address",
    ),
    "address_line_2": ("address 2", "address line 2", "suite", "unit"),
    "city": ("town",),
    "county": ("parish",),
    "state": ("province", "state province", "state/province"),
    "postal_code": ("zip", "zip code", "zipcode", "postal", "postal code"),
    "country": (),
    "latitude": ("lat",),
    "longitude": (
        "lon",
        "lng",
    ),
    "occupancy_type": ("occupancy",),
    "occupancy_subtype": ("occupancy sub type", "occupancy subtype"),
    "building_type": (),
    "construction_type": (),
    "roof_type": (),
    "foundation_type": (),
    "year_built": ("built year",),
    "year_renovated": ("renovated year",),
    "stories": ("story", "floors", "floor count"),
    "square_feet": ("sqft", "sq ft", "square ft", "area"),
    "sprinklered": ("sprinkler", "sprinklers"),
    "alarm_type": ("alarm",),
    "protection_class": (),
    "total_insured_value": ("tiv", "total value"),
    "building_value": (),
    "contents_value": (),
    "business_interruption_value": ("bi value", "business interruption"),
    "deductible": (),
    "flood_zone": (),
    "earthquake_zone": ("eq zone",),
    "hurricane_zone": (),
    "tornado_zone": (),
    "wildfire_zone": (),
    "hail_zone": (),
    "wind_zone": (),
    "storm_surge_zone": ("storm surge",),
    "tsunami_zone": (),
    "volcanic_zone": (),
    "landslide_zone": (),
    "sinkhole_zone": (),
    "winter_storm_zone": ("winter zone",),
    "lightning_zone": (),
    "climate_zone": (),
    "distance_to_coast_miles": ("distance to coast", "distance coast miles"),
    "distance_to_fault_miles": ("distance to fault", "distance fault miles"),
    "contact_name": ("contact",),
    "contact_email": ("email",),
    "contact_phone": ("phone", "telephone"),
}


def _normalize_header(value: Any) -> str:
    text = "" if value is None else str(value)
    return re.sub(r"[^a-z0-9]", "", text.strip().lower())


def _non_empty(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    return True


def _is_row_empty(row_values: Sequence[Any]) -> bool:
    return not any(_non_empty(value) for value in row_values)


def _location_model_fields() -> set[str]:
    return {field.name for field in fields(Location)}


def _extractable_fields() -> set[str]:
    return _location_model_fields() - _INTERNAL_FIELDS


def _alias_index() -> dict[str, set[str]]:
    aliases_by_field: dict[str, set[str]] = {}
    for field_name in sorted(_extractable_fields()):
        aliases = {
            field_name,
            field_name.replace("_", " "),
            field_name.replace("_", ""),
        }
        aliases.update(_FIELD_ALIASES.get(field_name, ()))
        aliases_by_field[field_name] = {_normalize_header(alias) for alias in aliases}
    return aliases_by_field


_ALIASES_BY_FIELD = _alias_index()


def _coerce_float(value: Any) -> float | None:
    if not _non_empty(value):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = str(value).strip().replace(",", "")
    if cleaned.endswith("%"):
        cleaned = cleaned[:-1]
    try:
        return float(cleaned)
    except ValueError:
        return None


def _coerce_int(value: Any) -> int | None:
    float_value = _coerce_float(value)
    if float_value is None:
        return None
    return int(float_value)


def _coerce_bool(value: Any) -> bool | None:
    if isinstance(value, bool):
        return value
    if not _non_empty(value):
        return None
    normalized = str(value).strip().lower()
    if normalized in {"true", "yes", "y", "1"}:
        return True
    if normalized in {"false", "no", "n", "0"}:
        return False
    return None


_FLOAT_FIELDS = {
    "latitude",
    "longitude",
    "square_feet",
    "total_insured_value",
    "building_value",
    "contents_value",
    "business_interruption_value",
    "deductible",
    "distance_to_coast_miles",
    "distance_to_fault_miles",
}
_INT_FIELDS = {"year_built", "year_renovated", "stories"}
_BOOL_FIELDS = {"sprinklered"}


def _coerce_field_value(field_name: str, value: Any) -> Any:
    if field_name in _FLOAT_FIELDS:
        return _coerce_float(value)
    if field_name in _INT_FIELDS:
        return _coerce_int(value)
    if field_name in _BOOL_FIELDS:
        return _coerce_bool(value)
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    return value


def _detect_header_row(
    rows: Sequence[Sequence[Any]],
) -> tuple[int, dict[int, str]] | None:
    best_row_index = -1
    best_mapping: dict[int, str] = {}
    best_score = 0

    for row_index, row_values in enumerate(rows[:_HEADER_SCAN_LIMIT]):
        mapping: dict[int, str] = {}
        already_mapped_fields: set[str] = set()

        for column_index, cell in enumerate(row_values):
            normalized = _normalize_header(cell)
            if not normalized:
                continue
            for field_name, aliases in _ALIASES_BY_FIELD.items():
                if normalized in aliases and field_name not in already_mapped_fields:
                    mapping[column_index] = field_name
                    already_mapped_fields.add(field_name)
                    break

        score = len(set(mapping.values()))
        has_address = "address_line_1" in mapping.values()
        if score > best_score and (has_address or score >= 3):
            best_row_index = row_index
            best_mapping = mapping
            best_score = score

    if best_row_index < 0 or best_score < 2:
        return None
    return best_row_index, best_mapping


def _build_location_from_row(
    row_values: Sequence[Any],
    column_mapping: Mapping[int, str],
    *,
    submission_id: str | None,
    source_file: str,
    source_sheet: str,
    source_row: int,
) -> Location | None:
    if _is_row_empty(row_values):
        return None

    parsed: dict[str, Any] = {}
    for column_index, field_name in column_mapping.items():
        value = row_values[column_index] if column_index < len(row_values) else None
        parsed[field_name] = _coerce_field_value(field_name, value)

    if not any(_non_empty(value) for value in parsed.values()):
        return None

    missing_critical = [
        field_name
        for field_name in _CRITICAL_FIELDS
        if not _non_empty(parsed.get(field_name))
    ]
    additional_fields: dict[str, Any] = {}
    if missing_critical:
        additional_fields["missing_critical_fields"] = missing_critical

    payload = {
        **parsed,
        "submission_id": submission_id,
        "source_file": source_file,
        "source_sheet": source_sheet,
        "source_row": source_row,
        "additional_fields": additional_fields,
    }
    return Location(**payload)


def _iter_excel_files(input_dir: Path) -> list[Path]:
    if not input_dir.exists() or not input_dir.is_dir():
        return []
    return sorted(
        path for path in input_dir.iterdir() if path.suffix.lower() in {".xlsx", ".xls"}
    )


def _iter_pdf_files(input_dir: Path) -> list[Path]:
    if not input_dir.exists() or not input_dir.is_dir():
        return []
    return sorted(path for path in input_dir.iterdir() if path.suffix.lower() == ".pdf")


def _iter_unsupported_files(input_dir: Path) -> list[Path]:
    if not input_dir.exists() or not input_dir.is_dir():
        return []
    supported = {".xlsx", ".xls", ".pdf"}
    return sorted(
        path
        for path in input_dir.iterdir()
        if path.is_file() and path.suffix.lower() not in supported
    )


def _split_files_by_type(
    files: Sequence[Path],
) -> tuple[list[Path], list[Path], list[Path]]:
    excel_files: list[Path] = []
    pdf_files: list[Path] = []
    unsupported_files: list[Path] = []
    for file_path in files:
        suffix = file_path.suffix.lower()
        if suffix in {".xlsx", ".xls"}:
            excel_files.append(file_path)
        elif suffix == ".pdf":
            pdf_files.append(file_path)
        else:
            unsupported_files.append(file_path)
    return (
        sorted(excel_files, key=lambda path: path.name),
        sorted(pdf_files, key=lambda path: path.name),
        sorted(unsupported_files, key=lambda path: path.name),
    )


def _extract_pdf_text_from_file(file_path: Path) -> str:
    if PdfReader is None:
        raise RuntimeError("PDF parsing requires optional dependency 'pypdf'")
    reader = PdfReader(file_path.as_posix())
    chunks: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        stripped = text.strip()
        if stripped:
            chunks.append(stripped)
    return "\n".join(chunks)


def _normalize_address_part(value: object | None) -> str:
    if value is None:
        return ""
    return re.sub(r"[^a-z0-9]", "", str(value).lower())


def _address_key(
    *,
    address_line_1: str | None,
    city: str | None,
    state: str | None,
    postal_code: str | None,
    country: str | None,
) -> str:
    if not _non_empty(address_line_1):
        return ""
    parts = [address_line_1, city, state, postal_code, country]
    return "|".join(_normalize_address_part(part) for part in parts)


def _location_address_key(location: Location) -> str:
    return _address_key(
        address_line_1=location.address_line_1,
        city=location.city,
        state=location.state,
        postal_code=location.postal_code,
        country=location.country,
    )


def _to_list(value: Any) -> list[Any]:
    if value is None:
        return []
    if isinstance(value, list):
        return list(value)
    if isinstance(value, tuple):
        return list(value)
    return [value]


@dataclass(frozen=True)
class ExtractedLocationCandidate:
    location_name: str | None
    address_line_1: str | None
    city: str | None
    state: str | None
    postal_code: str | None
    country: str | None
    confidence: float
    supplementary_text: str | None
    extraction_method: str


class PdfLocationExtractor(Protocol):
    def extract_location_candidates(
        self, text: str
    ) -> list[ExtractedLocationCandidate]: ...


_FALLBACK_LOCATION_PATTERN = re.compile(
    r"(?P<address>\d{1,6}\s+[A-Za-z0-9 .#'/-]{3,}?),\s*"
    r"(?P<city>[A-Za-z .'-]{2,}),\s*(?P<state>[A-Z]{2})\s+"
    r"(?P<postal>\d{5}(?:-\d{4})?)"
)


def _coerce_confidence(value: Any, *, default: float) -> float:
    parsed = _coerce_float(value)
    if parsed is None:
        return default
    bounded = max(0.0, min(1.0, parsed))
    return round(bounded, 4)


def _candidate_from_mapping(
    mapping: Mapping[str, Any], *, extraction_method: str
) -> ExtractedLocationCandidate:
    return ExtractedLocationCandidate(
        location_name=_coerce_field_value(
            "location_name", mapping.get("location_name")
        ),
        address_line_1=_coerce_field_value(
            "address_line_1", mapping.get("address_line_1")
        ),
        city=_coerce_field_value("city", mapping.get("city")),
        state=_coerce_field_value("state", mapping.get("state")),
        postal_code=_coerce_field_value("postal_code", mapping.get("postal_code")),
        country=_coerce_field_value("country", mapping.get("country")),
        confidence=_coerce_confidence(mapping.get("confidence"), default=0.8),
        supplementary_text=_coerce_field_value(
            "location_name", mapping.get("supplementary_text")
        ),
        extraction_method=extraction_method,
    )


def _fallback_extract_location_candidates(
    text: str,
) -> list[ExtractedLocationCandidate]:
    candidates: list[ExtractedLocationCandidate] = []
    for match in _FALLBACK_LOCATION_PATTERN.finditer(text):
        snippet_start = max(0, match.start() - 40)
        snippet_end = min(len(text), match.end() + 80)
        snippet = text[snippet_start:snippet_end].strip()
        candidates.append(
            ExtractedLocationCandidate(
                location_name=None,
                address_line_1=match.group("address").strip(),
                city=match.group("city").strip(),
                state=match.group("state").strip(),
                postal_code=match.group("postal").strip(),
                country=None,
                confidence=0.55,
                supplementary_text=snippet,
                extraction_method="fallback_regex",
            )
        )
    return candidates


class HybridPdfLocationExtractor:
    def __init__(self, llm_client: AzureOpenAIClient | None = None) -> None:
        self._llm_client = llm_client

    def extract_location_candidates(
        self, text: str
    ) -> list[ExtractedLocationCandidate]:
        if not text.strip():
            return []

        llm_candidates: list[ExtractedLocationCandidate] = []
        if self._llm_client is not None:
            raw = self._llm_client.extract_location_candidates(text)
            llm_candidates = [
                _candidate_from_mapping(entry, extraction_method="llm")
                for entry in raw
                if isinstance(entry, Mapping)
            ]
        if llm_candidates:
            return llm_candidates
        return _fallback_extract_location_candidates(text)


def _merge_pdf_metadata(
    location: Location,
    *,
    source_file: str,
    confidence: float,
    supplementary_text: str | None,
    extraction_method: str,
) -> Location:
    additional = dict(location.additional_fields)

    sources = _to_list(additional.get("pdf_sources"))
    if source_file not in sources:
        sources.append(source_file)
    additional["pdf_sources"] = sources

    confidence_values = _to_list(additional.get("pdf_extraction_confidence"))
    if confidence not in confidence_values:
        confidence_values.append(confidence)
    additional["pdf_extraction_confidence"] = confidence_values

    methods = _to_list(additional.get("pdf_extraction_methods"))
    if extraction_method not in methods:
        methods.append(extraction_method)
    additional["pdf_extraction_methods"] = methods

    if supplementary_text and supplementary_text not in _to_list(
        additional.get("pdf_supplementary")
    ):
        supplemental = _to_list(additional.get("pdf_supplementary"))
        supplemental.append(supplementary_text)
        additional["pdf_supplementary"] = supplemental

    if confidence < _LOW_CONFIDENCE_THRESHOLD:
        flags = _to_list(additional.get("audit_flags"))
        if "low_confidence_pdf_extraction" not in flags:
            flags.append("low_confidence_pdf_extraction")
        additional["audit_flags"] = flags

    return replace(location, additional_fields=additional)


def _build_location_from_pdf_candidate(
    candidate: ExtractedLocationCandidate,
    *,
    submission_id: str | None,
    source_file: str,
    source_row: int,
) -> Location | None:
    if not _non_empty(candidate.address_line_1):
        return None

    additional_fields: dict[str, Any] = {
        "pdf_sources": [source_file],
        "pdf_extraction_confidence": [candidate.confidence],
        "pdf_extraction_methods": [candidate.extraction_method],
    }
    if candidate.supplementary_text:
        additional_fields["pdf_supplementary"] = [candidate.supplementary_text]
    if candidate.confidence < _LOW_CONFIDENCE_THRESHOLD:
        additional_fields["audit_flags"] = ["low_confidence_pdf_extraction"]

    return Location(
        submission_id=submission_id,
        source_file=source_file,
        source_sheet="pdf",
        source_row=source_row,
        location_name=candidate.location_name,
        address_line_1=candidate.address_line_1,
        city=candidate.city,
        state=candidate.state,
        postal_code=candidate.postal_code,
        country=candidate.country,
        additional_fields=additional_fields,
    )


def _load_xlsx_with_retry(file_path: Path) -> Any:
    for attempt in range(1, _FILE_READ_RETRIES + 1):
        try:
            return load_workbook(file_path, read_only=True, data_only=True)
        except (PermissionError, OSError):
            if attempt == _FILE_READ_RETRIES:
                raise
            time.sleep(_RETRY_BACKOFF_SECONDS * attempt)


def _parse_xlsx_file(file_path: Path, submission_id: str | None) -> list[Location]:
    workbook = _load_xlsx_with_retry(file_path)
    locations: list[Location] = []
    try:
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            rows = list(worksheet.iter_rows(values_only=True))
            detected = _detect_header_row(rows)
            if detected is None:
                continue

            header_row_index, column_mapping = detected
            for offset, row_values in enumerate(
                rows[header_row_index + 1 :], start=header_row_index + 2
            ):
                location = _build_location_from_row(
                    row_values,
                    column_mapping,
                    submission_id=submission_id,
                    source_file=file_path.name,
                    source_sheet=sheet_name,
                    source_row=offset,
                )
                if location is not None:
                    locations.append(location)
    finally:
        workbook.close()
    return locations


def _parse_xls_file(file_path: Path, submission_id: str | None) -> list[Location]:
    if xlrd is None:
        # .xls parsing is optional because modern environments frequently omit xlrd.
        raise RuntimeError(".xls parsing requires optional dependency 'xlrd'")

    for attempt in range(1, _FILE_READ_RETRIES + 1):
        try:
            workbook = xlrd.open_workbook(file_path.as_posix(), on_demand=True)
            break
        except (PermissionError, OSError):
            if attempt == _FILE_READ_RETRIES:
                raise
            time.sleep(_RETRY_BACKOFF_SECONDS * attempt)
    else:  # pragma: no cover
        return []

    locations: list[Location] = []
    try:
        for sheet in workbook.sheets():
            rows = [sheet.row_values(index) for index in range(sheet.nrows)]
            detected = _detect_header_row(rows)
            if detected is None:
                continue

            header_row_index, column_mapping = detected
            for offset, row_values in enumerate(
                rows[header_row_index + 1 :], start=header_row_index + 2
            ):
                location = _build_location_from_row(
                    row_values,
                    column_mapping,
                    submission_id=submission_id,
                    source_file=file_path.name,
                    source_sheet=sheet.name,
                    source_row=offset,
                )
                if location is not None:
                    locations.append(location)
    finally:
        workbook.release_resources()
    return locations


class IngestionAgent(NoOpAgent):
    def __init__(
        self,
        config: AppConfig | None = None,
        pdf_location_extractor: PdfLocationExtractor | None = None,
    ) -> None:
        super().__init__("ingestion")
        self._config = config or AppConfig.from_env()
        self._pdf_location_extractor = (
            pdf_location_extractor
            or HybridPdfLocationExtractor(AzureOpenAIClient(self._config))
        )

    def parse_locations(
        self,
        submission_id: str | None = None,
        input_files: Sequence[str] | None = None,
    ) -> tuple[LocationManifest, list[str]]:
        if input_files is None:
            files = _iter_excel_files(self._config.input_dir)
            pdf_files = _iter_pdf_files(self._config.input_dir)
            unsupported_files = _iter_unsupported_files(self._config.input_dir)
        else:
            resolved_files: list[Path] = []
            for item in input_files:
                candidate = Path(item)
                if not candidate.is_absolute():
                    candidate = self._config.input_dir / candidate
                resolved_files.append(candidate)
            files, pdf_files, unsupported_files = _split_files_by_type(resolved_files)

        warnings: list[str] = []
        all_locations: list[Location] = []

        for file_path in unsupported_files:
            warning = f"Skipped unsupported file type '{file_path.name}'"
            warnings.append(warning)
            LOGGER.warning(warning)

        for file_path in files:
            try:
                if file_path.suffix.lower() == ".xlsx":
                    extracted = _parse_xlsx_file(file_path, submission_id)
                else:
                    extracted = _parse_xls_file(file_path, submission_id)
                all_locations.extend(extracted)
            except (  # continue processing remaining files by design
                RuntimeError,
                OSError,
                ValueError,
                TypeError,
                InvalidFileException,
                BadZipFile,
            ) as exc:
                warning = f"Skipped unreadable file '{file_path.name}': {exc}"
                warnings.append(warning)
                LOGGER.warning(warning)

        address_index: dict[str, int] = {}
        for index, location in enumerate(all_locations):
            key = _location_address_key(location)
            if key and key not in address_index:
                address_index[key] = index

        for file_path in pdf_files:
            try:
                text = _extract_pdf_text_from_file(file_path)
            except (RuntimeError, OSError, ValueError, TypeError) as exc:
                warning = f"Skipped unreadable file '{file_path.name}': {exc}"
                warnings.append(warning)
                LOGGER.warning(warning)
                continue

            candidates = self._pdf_location_extractor.extract_location_candidates(text)
            if not candidates:
                continue

            for row_number, candidate in enumerate(candidates, start=1):
                key = _address_key(
                    address_line_1=candidate.address_line_1,
                    city=candidate.city,
                    state=candidate.state,
                    postal_code=candidate.postal_code,
                    country=candidate.country,
                )
                if not key:
                    continue

                existing_index = address_index.get(key)
                if existing_index is not None:
                    all_locations[existing_index] = _merge_pdf_metadata(
                        all_locations[existing_index],
                        source_file=file_path.name,
                        confidence=candidate.confidence,
                        supplementary_text=candidate.supplementary_text,
                        extraction_method=candidate.extraction_method,
                    )
                    continue

                merged = _build_location_from_pdf_candidate(
                    candidate,
                    submission_id=submission_id,
                    source_file=file_path.name,
                    source_row=row_number,
                )
                if merged is None:
                    continue
                all_locations.append(merged)
                address_index[key] = len(all_locations) - 1

        manifest = LocationManifest.from_records(
            all_locations, submission_id=submission_id
        )
        return manifest, warnings

    def invoke(self, payload: Mapping[str, Any]) -> AgentResult:
        submission_id = payload.get("submission_id")
        input_files = payload.get("input_files")
        normalized_input_files: list[str] | None = None
        if isinstance(input_files, list):
            normalized_input_files = [str(item) for item in input_files]

        manifest, warnings = self.parse_locations(
            submission_id=submission_id,
            input_files=normalized_input_files,
        )
        return AgentResult(
            agent=self.name,
            status="ok",
            payload={
                "received": dict(payload),
                "message": "ingestion completed",
                "location_count": manifest.count,
                "manifest_hash": manifest.manifest_hash,
                "manifest": manifest.to_dict(),
                "warnings": warnings,
                "input_files": normalized_input_files,
            },
        )
